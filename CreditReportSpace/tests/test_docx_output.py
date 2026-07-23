"""DOCX 產出測試：以去識別化 fixture 生成，驗證表格結構與內容。"""

import copy

import pytest
from docx import Document

from credit_report import build_docx, validate


@pytest.fixture
def built_doc(sample_data, tmp_path):
    out = tmp_path / "report.docx"
    build_docx(sample_data, out)
    return Document(str(out))


def _cell_texts(table):
    return [[c.text for c in row.cells] for row in table.rows]


def test_two_tables_generated(built_doc):
    assert len(built_doc.tables) == 2


def test_header_and_titles(built_doc):
    texts = [p.text for p in built_doc.paragraphs]
    assert any("伍、金融借款：(仟元)" in t and "資料日期-2026年6月底" in t for t in texts)
    assert any(t.startswith("◎虛構測試建設有限公司（TEST）") for t in texts)
    assert any(t.startswith("◎負責人-測試甲（虛構人名）") for t in texts)


def test_simple_table_shape_and_totals(built_doc):
    t = built_doc.tables[0]
    rows = _cell_texts(t)
    assert rows[0] == ["主/從債", "往來銀行", "項目", "借款金額", "借款餘額", "備註"]
    # 5 資料列 + 表頭 + 合計
    assert len(rows) == 7
    total = rows[-1]
    assert "合計" in total[0]
    assert "200,500" in total
    assert "198,700" in total


def test_extended_table_headers(built_doc):
    t = built_doc.tables[1]
    head = _cell_texts(t)[0]
    assert head[0] == "主/從"
    assert "前次餘額\n2023" in head
    assert "前次餘額\n2020.06" in head
    assert "增減情形" in head


def test_role_repeated_on_every_data_row(built_doc):
    """主/從欄逐列重複（不合併），任何分頁位置都看得到角色。"""
    t = built_doc.tables[0]
    for r in range(1, 6):  # 5 資料列
        assert t.rows[r].cells[0].text == "主債"
    # 不是垂直合併：相鄰列是不同儲存格
    assert t.rows[1].cells[0]._tc is not t.rows[2].cells[0]._tc


def test_bank_cells_merged_for_consecutive_same_bank(built_doc):
    t = built_doc.tables[0]
    # 華銀/測試北 兩列 → 銀行儲存格合併
    assert t.rows[1].cells[1]._tc is t.rows[2].cells[1]._tc
    assert t.rows[1].cells[1].text == "華銀/測試北"


def test_negative_change_rendered_with_parens(built_doc):
    t = built_doc.tables[1]
    all_text = "\n".join("\t".join(c.text for c in r.cells) for r in t.rows)
    assert "(50)" in all_text
    assert "(700)" in all_text


def test_struck_row_has_strikethrough_runs(built_doc):
    t = built_doc.tables[1]
    struck_runs = [
        run
        for row in t.rows
        for cell in row.cells
        for p in cell.paragraphs
        for run in p.runs
        if run.font.strike
    ]
    assert any("轉貸至他行" in r.text for r in struck_runs)
    assert any("80,000" == r.text for r in struck_runs)


def test_needs_review_annotation(built_doc):
    t = built_doc.tables[1]
    all_text = "\n".join(c.text for r in t.rows for c in r.cells)
    assert "（待人工確認）" in all_text


def test_na_rendered_as_dash(built_doc):
    """prev_balances 的 \"N/A\" 渲染為「—」。"""
    t = built_doc.tables[1]
    for row in t.rows:
        cells = [c.text for c in row.cells]
        if any("土銀/測試五" in c for c in cells):
            assert "—" in cells
            break
    else:
        pytest.fail("找不到 N/A 測試列")


def test_numeric_cells_have_no_wrap(built_doc):
    """數字欄有 w:noWrap，避免括號金額折行。"""
    from docx.oxml.ns import qn

    t = built_doc.tables[1]
    data_cell = t.rows[1].cells[3]
    assert data_cell._tc.get_or_add_tcPr().find(qn("w:noWrap")) is not None
    total_change_cell = t.rows[-1].cells[7]
    assert total_change_cell._tc.get_or_add_tcPr().find(qn("w:noWrap")) is not None


def test_notes_and_source(built_doc):
    texts = [p.text for p in built_doc.paragraphs]
    assert any("說明：" in t for t in texts)
    assert any("申戶銀行往來四家" in t for t in texts)
    assert any("資料來源：財團法人金融聯合徵信中心信用報告。" in t for t in texts)


def test_null_amount_rendered_blank(built_doc):
    """struck 列 balance=null → 餘額欄空白。"""
    t = built_doc.tables[1]
    for row in t.rows:
        cells = [c.text for c in row.cells]
        if any("轉貸至他行" in c for c in cells):
            assert cells[4] == ""
            break
    else:
        pytest.fail("找不到轉貸列")


def test_group_without_total_renders(sample_data, tmp_path):
    data = copy.deepcopy(sample_data)
    del data["entities"][0]["groups"][0]["total"]
    assert validate(data).ok
    out = tmp_path / "no_total.docx"
    build_docx(data, out)
    t = Document(str(out)).tables[0]
    assert len(t.rows) == 6  # 表頭 + 5 資料列，無合計列
