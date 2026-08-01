"""空白模板 PII 守衛：模板為二進位 docx，不受 test_pii_lint 文字掃描涵蓋，
故單獨開檔掃全部文字。PII 樣式與真實樣本雜湊自 test_pii_lint 匯入（唯一存放處，
避免雜湊 hex 的長數字串在他處觸發 lint）。"""

import re
from pathlib import Path

from docx import Document

from test_pii_lint import TW_ID, LONG_DIGITS, _hash_hits

SPACE_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = SPACE_ROOT / "templates" / "credit_report_blank.docx"


def _all_text():
    doc = Document(str(TEMPLATE))
    parts = [p.text for p in doc.paragraphs]
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_template_exists():
    assert TEMPLATE.is_file()


def test_template_has_no_pii():
    text = _all_text()
    assert not TW_ID.findall(text), "模板殘留身分證樣式"
    assert not LONG_DIGITS.findall(text), "模板殘留 9+ 位長數字"
    assert not _hash_hits(text), "模板命中真實樣本雜湊"


def test_template_structure_intact():
    doc = Document(str(TEMPLATE))
    assert len(doc.tables) == 21
    titles = "\n".join(p.text for p in doc.paragraphs)
    for seg in ["壹、", "肆、", "陸、金融借款", "柒、", "捌、保證人", "玖、"]:
        assert seg in titles
