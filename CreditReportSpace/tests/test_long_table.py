"""長表（42+ 列）回歸：主/從欄逐列重複、不做跨頁合併；
本機有 LibreOffice + Poppler 時做真渲染，並逐頁以 PDF 文字斷言角色可見。"""

import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from credit_report import build_docx, validate
import render_docx

N_ROWS = 45


def _long_report(row_count=N_ROWS):
    rows = [
        {
            "bank": f"測試{i // 3}銀/分行{i}",
            "item": ("中擔", "短擔", "長擔")[i % 3],
            "loan_amount": 1000 + i,
            "balance": 900 + i,
        }
        for i in range(row_count)
    ]
    total = {
        "loan_amount": sum(r["loan_amount"] for r in rows),
        "balance": sum(r["balance"] for r in rows),
    }
    return {
        "meta": {"data_date": "資料日期-2026年6月底（TEST 長表）"},
        "entities": [{
            "name": "長表測試建設有限公司（虛構）",
            "layout": "simple",
            "groups": [{"role": "主債", "rows": rows, "total": total}],
        }],
    }


@pytest.fixture(scope="module")
def long_docx(tmp_path_factory):
    data = _long_report()
    assert validate(data).ok
    out = tmp_path_factory.mktemp("long") / "long_table.docx"
    build_docx(data, out)
    return out


def test_role_on_every_data_row_without_merge(long_docx):
    from docx.oxml.ns import qn

    table = Document(str(long_docx)).tables[0]
    data_rows = table.rows[1:1 + N_ROWS]
    for row in data_rows:
        assert row.cells[0].text == "主債"
        # 逐列獨立儲存格：不得有垂直合併標記（合併的續列 text 會是空白且
        # 帶 w:vMerge，跨頁時整頁看不到角色）
        tcpr = row.cells[0]._tc.find(qn("w:tcPr"))
        assert tcpr is None or tcpr.find(qn("w:vMerge")) is None


def test_long_table_total_row_present(long_docx):
    table = Document(str(long_docx)).tables[0]
    assert len(table.rows) == 1 + N_ROWS + 1
    assert "合計" in table.rows[-1].cells[0].text


needs_render_tools = pytest.mark.skipif(
    render_docx.find_soffice() is None
    or shutil.which("pdftoppm") is None
    or shutil.which("pdftotext") is None,
    reason="需要本機 LibreOffice + Poppler（pdftoppm/pdftotext）",
)


@needs_render_tools
def test_long_table_every_page_shows_header_and_role(long_docx, tmp_path):
    pages = render_docx.render(Path(long_docx), tmp_path)
    assert len(pages) >= 2

    pdf = tmp_path / long_docx.stem / (long_docx.stem + ".pdf")
    assert pdf.is_file()
    text = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"],
        check=True, capture_output=True, timeout=120,
    ).stdout.decode("utf-8")
    page_texts = [p for p in text.split("\f") if p.strip()]
    assert len(page_texts) == len(pages)
    for number, page in enumerate(page_texts, start=1):
        # 每個資料列恰有一個項目簡稱（中擔/短擔/長擔），以此當資料列計數；
        # 銀行名含 ASCII 會被斷行，不適合當標記。
        data_row_marks = sum(page.count(k) for k in ("中擔", "短擔", "長擔"))
        if data_row_marks:
            # 逐頁斷言：重複表頭與每一資料列的角色都看得到
            assert "往來銀行" in page, f"第 {number} 頁缺表頭"
            assert page.count("主債") >= data_row_marks, (
                f"第 {number} 頁有資料列缺主/從角色標示"
                f"（主債×{page.count('主債')} < 資料列×{data_row_marks}）"
            )
