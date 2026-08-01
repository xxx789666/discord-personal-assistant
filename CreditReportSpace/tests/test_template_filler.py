"""填格引擎測試：以去識別化（虛構）段落資料填空白模板，
驗證陸段插入、損益計算（四捨五入基本項→算衍生列＋暫結慣例）、
捌保證人欄與 needs_review 淡黃底。"""

import pytest
from docx import Document as _Doc
from docx import Document
from docx.oxml.ns import qn

from credit_report import template_filler as TF
from credit_report.template_filler import (
    fill_report, _compute_income_year, _fill_luduan, _fill_income,
)


# ── 虛構 fixtures ─────────────────────────────────────────
LUDUAN = {
    "meta": {"section_title": "陸、金融借款：（仟元）",
             "data_date": "資料日期-115年3月底（虛構）", "source_note": ""},
    "entities": [{
        "name": "測試甲（虛構）", "layout": "simple",
        "groups": [
            {"role": "主債", "rows": [
                {"bank": "測試銀行/測試北", "item": "中擔",
                 "loan_amount": 1000, "balance": 900, "note": "測試"},
                {"bank": "測試銀行/測試北", "item": "長擔",
                 "loan_amount": 500, "balance": 480},
            ], "total": {"balance": 1380}},
            {"role": "從債", "rows": [
                {"bank": "測試商銀/測試南", "item": "中放",
                 "balance": 200, "note": "測試乙"},
            ]},
        ],
        "notes": ["測試主債餘額 1,380 仟元（虛構）。"],
    }],
}

INCOME = {"years": [
    {"label": "2099(暫結)", "provisional": True, "sales_net": 1000,
     "cogs": 600, "opex": 200, "other_income": 10, "interest_exp": 30,
     "other_exp": None, "tax": 50},
    {"label": "2098 年", "provisional": False, "sales_net": 800,
     "cogs": 500, "opex": 150, "other_income": 5, "interest_exp": 20,
     "other_exp": None, "tax": 40},
    {"label": "2097 年", "missing": True},
]}

GUARANTOR = {"guarantors": [
    {"name": "測試甲", "birth": "80.01.01", "gender": "男",
     "id_no": "測試證號X", "address": "測試市測試區測試路1號",
     "married": "已婚", "job": None, "phone": None},
]}


def _cell_fill(cell):
    tcpr = cell._tc.tcPr
    if tcpr is None:
        return None
    shd = tcpr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def _table_by_header(doc, first):
    for tb in doc.tables:
        if tb.rows[0].cells[0].text.strip().startswith(first):
            return tb
    return None


@pytest.fixture
def built(tmp_path):
    out = tmp_path / "r.docx"
    _, review = fill_report(
        {"luduan": LUDUAN, "income": INCOME, "guarantor": GUARANTOR}, out)
    return Document(str(out)), review


# ── 損益計算（純函式）─────────────────────────────────────
def test_income_provisional_uses_pretax_as_aftertax():
    d = _compute_income_year(INCOME["years"][0])
    assert d["gross"] == 400          # 1000-600
    assert d["op_income"] == 200      # 400-200
    assert d["pretax"] == 180         # 200+10-30
    assert d["aftertax"] == 180       # 暫結：=稅前
    assert d["tax"] is None           # 暫結：所得稅留空


def test_income_final_year_deducts_tax():
    d = _compute_income_year(INCOME["years"][1])
    assert d["gross"] == 300
    assert d["op_income"] == 150
    assert d["pretax"] == 135         # 150+5-20
    assert d["aftertax"] == 95        # 135-40
    assert d["tax"] == 40


# ── 損益表填入 ───────────────────────────────────────────
def test_income_table_values_and_percent(built):
    doc, _ = built
    tb = _table_by_header(doc, "項目/年度")
    rows = {r.cells[0].text.strip().lstrip("－-").strip(): r for r in tb.rows[1:]}
    # 2099 暫結欄（值 col1, % col2）
    assert rows["銷貨淨額"].cells[1].text == "1,000"
    assert rows["銷貨淨額"].cells[2].text == "100.00"
    assert rows["銷貨毛利"].cells[1].text == "400"
    assert rows["稅後淨利"].cells[1].text == "180"
    assert rows["所得稅"].cells[1].text == ""       # 暫結留空
    # 2098 確定欄（值 col3）
    assert rows["稅後淨利"].cells[3].text == "95"
    assert rows["所得稅"].cells[3].text == "40"


def test_income_missing_year_flagged(built):
    doc, review = built
    tb = _table_by_header(doc, "項目/年度")
    # 第三年（col5/col6）整欄空白且淡黃底
    net = [r for r in tb.rows[1:] if r.cells[0].text.strip() == "銷貨淨額"][0]
    assert net.cells[5].text == ""
    assert _cell_fill(net.cells[5]) == "FFF2CC"
    assert any("2097" in r and "缺" in r for r in review)


# ── 陸段插入 ─────────────────────────────────────────────
def test_luduan_inserted_with_date_header(built):
    doc, _ = built
    tb = _table_by_header(doc, "主/從債")
    assert tb.rows[0].cells[4].text == "2026/03"   # 由 data_date 推
    # 主債兩列 + 合計 + 從債一列
    texts = [[c.text for c in r.cells] for r in tb.rows]
    assert any("合計" in row[0] and "1,380" in row for row in texts)
    assert any("測試乙" in row[-1] for row in texts)


def test_luduan_preserves_following_section(built):
    doc, _ = built
    titles = "\n".join(p.text for p in doc.paragraphs)
    assert "柒、" in titles           # 下一段標題未被刪
    assert "◎測試甲（虛構）" in titles


# ── 捌 保證人 ────────────────────────────────────────────
def test_guarantor_basic_fields_filled(built):
    doc, _ = built
    tb = [t for t in doc.tables
          if [c.text for c in t.rows[0].cells][:5]
          == ["姓名", "出生年月日", "婚姻", "現職", "電話"]][0]
    assert tb.rows[1].cells[0].text == "測試甲"
    assert tb.rows[1].cells[2].text == "已婚"
    assert tb.rows[3].cells[0].text == "男"
    assert tb.rows[3].cells[1].text == "測試證號X"
    assert "測試路1號" in tb.rows[3].cells[2].text


def test_guarantor_missing_job_phone_flagged(built):
    doc, review = built
    tb = [t for t in doc.tables
          if [c.text for c in t.rows[0].cells][:5]
          == ["姓名", "出生年月日", "婚姻", "現職", "電話"]][0]
    assert tb.rows[1].cells[3].text == ""             # 現職空
    assert _cell_fill(tb.rows[1].cells[3]) == "FFF2CC"  # 淡黃底
    assert _cell_fill(tb.rows[1].cells[4]) == "FFF2CC"  # 電話淡黃底
    assert any("現職/電話" in r for r in review)


def test_guarantor_principal_ordered_first(tmp_path):
    """負責人（principal=true）排（1），即使在資料中列於後。"""
    section = {"guarantors": [
        {"name": "測試乙", "birth": "70.05.05", "gender": "女",
         "id_no": "測試證號Y", "address": "測試市A", "married": "已婚",
         "principal": False},
        {"name": "測試甲", "birth": "80.01.01", "gender": "男",
         "id_no": "測試證號X", "address": "測試市B", "married": "已婚",
         "principal": True},
    ]}
    out = tmp_path / "g.docx"
    fill_report({"guarantor": section}, out)
    doc = Document(str(out))
    gtabs = [t for t in doc.tables
             if [c.text for c in t.rows[0].cells][:5]
             == ["姓名", "出生年月日", "婚姻", "現職", "電話"]]
    assert gtabs[0].rows[1].cells[0].text == "測試甲"   # 負責人在（1）
    assert gtabs[1].rows[1].cells[0].text == "測試乙"


def test_balance_header_parsing():
    assert TF._luduan_balance_header("資料日期-115年3月底") == "2026/03"
    assert TF._luduan_balance_header("資料日期-2024年1月底") == "2024/01"
    assert TF._luduan_balance_header("無日期") == "借款餘額"


# ── 邊界與錯誤路徑 ───────────────────────────────────────
def test_luduan_missing_anchor_raises():
    doc = _Doc()  # 空白文件，無「陸、金融借款」標題
    doc.add_paragraph("無關內容")
    with pytest.raises(ValueError, match="陸、金融借款"):
        _fill_luduan(doc, LUDUAN)


def test_income_missing_table_raises():
    doc = _Doc()  # 無損益表
    with pytest.raises(ValueError, match="損益表"):
        _fill_income(doc, INCOME)


def test_luduan_needs_review_row_shaded(tmp_path):
    section = {
        "meta": {"data_date": "資料日期-115年3月底", "source_note": ""},
        "entities": [{"name": "測試（虛構）", "layout": "simple", "groups": [
            {"role": "主債", "rows": [
                {"bank": "測試銀行/測試", "item": "中擔", "loan_amount": 100,
                 "balance": None, "needs_review": True, "note": "模糊"}]}]}],
    }
    out = tmp_path / "r.docx"
    fill_report({"luduan": section}, out)
    doc = Document(str(out))
    tb = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "主/從債"][0]
    # 資料列每格淡黃底，備註含「待人工確認」
    row = tb.rows[1]
    assert _cell_fill(row.cells[0]) == "FFF2CC"
    assert "待人工確認" in row.cells[5].text


def test_luduan_struck_row_renders(tmp_path):
    section = {
        "meta": {"data_date": "資料日期-115年3月底", "source_note": ""},
        "entities": [{"name": "測試（虛構）", "layout": "simple", "groups": [
            {"role": "主債", "rows": [
                {"bank": "測試銀行/測試", "item": "中擔", "loan_amount": 100,
                 "balance": None, "struck": True, "note": "已結清轉貸"}]}]}],
    }
    out = tmp_path / "r.docx"
    _, review = fill_report({"luduan": section}, out)
    doc = Document(str(out))
    tb = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "主/從債"][0]
    cell = tb.rows[1].cells[2]  # 項目格
    strike = cell.paragraphs[0].runs[0].font.strike
    assert strike is True


def test_luduan_same_bank_rows_merged(built):
    doc, _ = built
    tb = _table_by_header(doc, "主/從債")
    # 主債兩列同銀行（測試銀行/測試北）→ 銀行儲存格垂直合併為同一 tc
    assert tb.rows[1].cells[1]._tc is tb.rows[2].cells[1]._tc


def test_guarantor_overflow_reported(tmp_path):
    # 模板僅 2 個保證人區塊；給 3 位 → 第三位未填並回報
    section = {"guarantors": [
        {"name": f"測試{i}", "birth": "80.01.01", "gender": "男",
         "id_no": f"測試證號{i}", "address": "測試市", "married": "已婚"}
        for i in range(3)]}
    out = tmp_path / "r.docx"
    _, review = fill_report({"guarantor": section}, out)
    assert any("超過模板" in r for r in review)
